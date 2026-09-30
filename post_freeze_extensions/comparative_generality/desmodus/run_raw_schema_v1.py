#!/usr/bin/env python3
from __future__ import annotations

import hashlib, io, json, os
from pathlib import Path

import requests
from openpyxl import load_workbook
from remotezip import RemoteZip

ROOT=Path(__file__).resolve().parents[3]
CONTRACT=ROOT/"post_freeze_extensions/comparative_generality/desmodus/raw_schema_contract_v1.json"
OUT=ROOT/"post_freeze_extensions/comparative_generality/desmodus/raw_schema_result_v1.json"
OUT_MD=ROOT/"post_freeze_extensions/comparative_generality/desmodus/RAW_SCHEMA_RESULT_V1.md"

def norm(x):
    return str(x).strip().lower().replace(" ","_").replace(".","_").replace("-","_")

def main():
    c=json.loads(CONTRACT.read_text())
    token=os.environ.get("DRYAD_API_TOKEN","").strip()
    if not token:
        raise RuntimeError("DRYAD_API_TOKEN missing")

    url=f"https://datadryad.org/api/v2/files/{c['archive']['dryad_file_id']}/download"
    rz=RemoteZip(url,headers={
        "Authorization":f"Bearer {token}",
        "User-Agent":"batter-desmodus-raw-schema-audit-v1/1.0"
    })
    info=rz.getinfo(c["archive"]["inner_path"])
    if int(info.file_size)!=int(c["archive"]["inner_uncompressed_bytes"]):
        raise RuntimeError("inner member size changed relative to frozen inventory")
    with rz.open(info) as f:
        raw=f.read()
    rz.close()
    sha=hashlib.sha256(raw).hexdigest()

    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
    sheets={}
    candidates=[]
    for ws in wb.worksheets:
        row=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
        headers=["" if x is None else str(x).strip() for x in row]
        sheets[ws.title]=headers
        n=[norm(x) for x in headers]
        has_id=any(x in n for x in ["individual","id_bat","id_bat_ring","id_tag","tag","band"])
        has_lon=any("long" in x or x in {"x","longitude"} for x in n)
        has_lat=any("lat" in x or x in {"y","latitude"} for x in n)
        has_alt=any("alt" in x for x in n)
        has_time=any(x in n for x in ["date_time","datetime","time","date","day","hour","utc_5"])
        if has_id and has_lon and has_lat and has_alt and has_time:
            candidates.append(ws.title)
    wb.close()

    payload={
        "schema_version":1,
        "study_id":c["study_id"],
        "gps_data_rows_read":False,
        "numeric_altitude_values_read":False,
        "inner_sha256":sha,
        "sheet_headers":sheets,
        "candidate_event_sheets":candidates,
        "decision":"SCHEMA_PASS" if len(candidates)==1 else "SCHEMA_STOP"
    }
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(payload,indent=2,sort_keys=True)+"\n",encoding="utf-8")

    lines=[
        "# Desmodus raw GPS schema audit v1","",
        "**Only workbook sheet names and row-1 headers were read. No GPS data row or altitude value was read.**","",
        f"- inner XLSX SHA256: `{sha}`",
        f"- sheets: **{len(sheets)}**",
        f"- candidate event sheets: **{candidates}**",""
    ]
    for name,headers in sheets.items():
        lines += [f"## {name}","",", ".join(f"`{x}`" for x in headers if x),""]

    lines += [f"Decision: **{payload['decision']}**",""]
    OUT_MD.write_text("\n".join(lines),encoding="utf-8")
    print(json.dumps({
        "inner_sha256":sha,
        "sheets":list(sheets),
        "candidate_event_sheets":candidates,
        "headers":sheets,
        "decision":payload["decision"]
    },sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
