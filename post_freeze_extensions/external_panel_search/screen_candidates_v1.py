#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
CONTRACT=ROOT/"post_freeze_extensions/external_panel_search/candidate_screen_contract_v1.json"
OUT=ROOT/"post_freeze_extensions/external_panel_search/candidate_screen_result_v1.json"
OUT_MD=ROOT/"post_freeze_extensions/external_panel_search/CANDIDATE_SCREEN_RESULT_V1.md"

def canon(x):
    return re.sub(r"_+","_",str(x).strip().lower().replace("-","_").replace(" ","_").replace(":","_")).strip("_")

def fetch(url):
    r=requests.get(url,headers={"User-Agent":"batter-external-screen-v1/1.0"},timeout=120)
    r.raise_for_status()
    return r.content

def finite_num(x):
    try:
        z=float(str(x).strip())
        return math.isfinite(z)
    except Exception:
        return False

def nonempty_string(x):
    # Intentionally never parse vertical magnitudes.
    if x is None:
        return False
    s=str(x).strip()
    return s!="" and s.lower() not in {"na","nan","none","null"}

def parse_dt(s):
    s=str(s).strip()
    fmts=[
        "%Y/%m/%d %H:%M:%S","%Y-%m-%d %H:%M:%S","%m/%d/%Y %H:%M:%S",
        "%Y/%m/%d %H:%M","%Y-%m-%d %H:%M","%m/%d/%Y %H:%M",
        "%Y/%m/%d","%Y-%m-%d","%m/%d/%Y"
    ]
    for fmt in fmts:
        try:
            return datetime.strptime(s,fmt),("%H" in fmt)
        except Exception:
            pass
    try:
        x=datetime.fromisoformat(s)
        return x,("T" in s or ":" in s)
    except Exception:
        raise ValueError(f"unparsed timestamp: {s!r}")

def session_label(raw,rule):
    dt,has_time=parse_dt(raw)
    if "date-only" in rule.lower() and not has_time:
        return dt.date().isoformat()
    if "shifted" in rule.lower():
        return (dt-timedelta(hours=12)).date().isoformat()
    return dt.date().isoformat()

def rows_csv(data):
    text=data.decode("utf-8-sig",errors="strict")
    rd=csv.DictReader(io.StringIO(text,newline=""))
    headers=rd.fieldnames or []
    return headers,[dict(r) for r in rd]

def rows_xlsx(data):
    tmp=Path("/tmp/batter_candidate.xlsx")
    tmp.write_bytes(data)
    wb=load_workbook(tmp,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True)
    headers=[str(x).strip() if x is not None else "" for x in next(it)]
    rows=[]
    for vals in it:
        rows.append({headers[i]:vals[i] for i in range(min(len(headers),len(vals)))})
    wb.close()
    tmp.unlink(missing_ok=True)
    return headers,rows

def resolve_header(headers,want):
    if want in headers:
        return want
    cmap={canon(h):h for h in headers}
    key=canon(want)
    if key not in cmap:
        raise RuntimeError(f"missing expected column {want!r}; headers={headers}")
    return cmap[key]

def screen_candidate(cand,gate):
    data=fetch(cand["file_url"])
    sha=hashlib.sha256(data).hexdigest()
    size=len(data)
    if cand["file_name"].lower().endswith(".csv"):
        headers,rows=rows_csv(data)
    elif cand["file_name"].lower().endswith((".xlsx",".xlsm")):
        headers,rows=rows_xlsx(data)
    else:
        raise RuntimeError(f"unsupported file type {cand['file_name']}")

    schema=cand["expected_schema"]
    hid=resolve_header(headers,schema["individual"])
    ht=resolve_header(headers,schema["timestamp"])
    hx=resolve_header(headers,schema["longitude"])
    hy=resolve_header(headers,schema["latitude"])
    hz=resolve_header(headers,schema["vertical"])

    per_ind_session=defaultdict(lambda:defaultdict(int))
    raw_ind=set()
    xyv_ind=set()
    valid_rows=0
    bad_time=0
    vertical_nonempty_rows=0

    for r in rows:
        iid=str(r.get(hid,"")).strip()
        if iid:
            raw_ind.add(iid)
        v_present=nonempty_string(r.get(hz))
        if v_present:
            vertical_nonempty_rows+=1
        if not iid or not finite_num(r.get(hx)) or not finite_num(r.get(hy)) or not v_present:
            continue
        try:
            sess=session_label(r.get(ht),cand["session_rule"])
        except Exception:
            bad_time+=1
            continue
        valid_rows+=1
        xyv_ind.add(iid)
        per_ind_session[iid][sess]+=1

    minfix=int(gate["minimum_fixes_per_session"])
    eligible_sessions={}
    repeat_ids=[]
    for iid in sorted(xyv_ind):
        sess={k:int(v) for k,v in sorted(per_ind_session[iid].items())}
        elig={k:v for k,v in sess.items() if v>=minfix}
        eligible_sessions[iid]={
            "session_count":len(sess),
            "eligible_session_count":len(elig),
            "session_fix_counts":sess,
            "eligible_session_fix_counts":elig
        }
        if len(elig)>=2:
            repeat_ids.append(iid)

    pass_n=len(xyv_ind)>=int(gate["minimum_individuals_with_xy_vertical_presence"])
    pass_rep=len(repeat_ids)>=int(gate["minimum_repeat_individuals"])
    passed=pass_n and pass_rep and vertical_nonempty_rows>0

    return {
        "candidate_id":cand["id"],
        "repository":cand["repository"],
        "doi":cand["doi"],
        "taxon":cand["taxon"],
        "file_name":cand["file_name"],
        "file_stream_id":cand["file_stream_id"],
        "sha256":sha,
        "size_bytes":size,
        "headers":headers,
        "vertical_field":hz,
        "numeric_vertical_values_read":False,
        "raw_row_count":len(rows),
        "raw_individual_count":len(raw_ind),
        "xy_vertical_presence_row_count":valid_rows,
        "xy_vertical_presence_individual_count":len(xyv_ind),
        "timestamp_parse_failures_among_xyv_rows":bad_time,
        "vertical_nonempty_row_count":vertical_nonempty_rows,
        "repeat_individual_count":len(repeat_ids),
        "repeat_individual_ids":repeat_ids,
        "individual_session_structure":eligible_sessions,
        "gate_checks":{
            "minimum_individuals_met":pass_n,
            "minimum_repeat_individuals_met":pass_rep,
            "minimum_fixes_per_session":minfix,
            "native_vertical_field_present":hz in headers
        },
        "decision":"PASS" if passed else "FAIL",
        "failure_reason":None if passed else (
            "fewer_than_8_individuals" if not pass_n else
            "fewer_than_5_repeat_individuals_with_two_50fix_sessions" if not pass_rep else
            "native_vertical_unavailable"
        )
    }

def main():
    c=json.loads(CONTRACT.read_text())
    results=[]
    for cand in c["candidates"]:
        results.append(screen_candidate(cand,c["structural_gate"]))

    payload={
        "schema_version":1,
        "study_id":c["study_id"],
        "parent_contract":c["parent_contract"],
        "numeric_vertical_values_read":False,
        "results":results,
        "passing_candidates":[r["candidate_id"] for r in results if r["decision"]=="PASS"],
        "search_outcome":"PASS_EXISTS" if any(r["decision"]=="PASS" for r in results) else "NO_PASS_IN_SCREENED_CANDIDATES"
    }
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(payload,indent=2,sort_keys=True)+"\n",encoding="utf-8")

    lines=[
        "# External Dryad candidate structural screen v1","",
        "**OUTCOME-BLIND: numeric vertical magnitudes were never parsed.**","",
        "| candidate | rows | xy+vertical individuals | repeat individuals | decision |",
        "|---|---:|---:|---:|---|"
    ]
    for r in results:
        lines.append(f"| {r['candidate_id']} | {r['raw_row_count']} | {r['xy_vertical_presence_individual_count']} | {r['repeat_individual_count']} | **{r['decision']}** |")
    lines += ["","Passing candidates: "+(", ".join(payload["passing_candidates"]) if payload["passing_candidates"] else "**none**"),""]
    OUT_MD.write_text("\n".join(lines),encoding="utf-8")

    print(json.dumps({
        "passing_candidates":payload["passing_candidates"],
        "results":{r["candidate_id"]:{
            "rows":r["raw_row_count"],
            "individuals":r["xy_vertical_presence_individual_count"],
            "repeat_individuals":r["repeat_individual_count"],
            "decision":r["decision"],
            "failure_reason":r["failure_reason"],
            "sha256":r["sha256"]
        } for r in results}
    },sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
