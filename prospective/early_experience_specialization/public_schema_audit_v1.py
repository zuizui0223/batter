#!/usr/bin/env python3
"""Schema-only audit of the public early-experience Mendeley ZIP.

Authorized:
- archive member names/sizes;
- CSV/TSV header only;
- XLSX sheet names/dimensions/header row only;
- MAT variable names/shapes/classes only;
- analysis-code identifiers / quoted column-like strings only.

No research data rows or numeric outcome arrays are reported.
"""

from __future__ import annotations
from pathlib import Path
import io, json, re, tempfile, urllib.request, zipfile, csv

HERE=Path(__file__).resolve().parent
OUT=HERE/"PUBLIC_SCHEMA_AUDIT_V1.json"
OUTMD=HERE/"PUBLIC_SCHEMA_AUDIT_V1.md"

ZIP_URL="https://data.mendeley.com/api/datasets-v2/datasets/wh7c636y3t/zip?version=1"
HEADERS={"User-Agent":"Mozilla/5.0 batter-early-experience-schema/1.0"}

RELEVANT=re.compile(
    r"(bat|id|season|trial|condition|enrich|impover|bold|risk|explor|activ|"
    r"colony|origin|sex|age|gps|distance|outdoor|area|personality)",
    re.I,
)

def fetch_bytes(url):
    req=urllib.request.Request(url,headers=HEADERS)
    with urllib.request.urlopen(req,timeout=180) as r:
        return r.read()

def csv_header(data,name):
    text=data.decode("utf-8-sig","replace")
    first=text.splitlines()[0] if text.splitlines() else ""
    delim="\t" if name.lower().endswith((".tsv",".tab")) else ","
    try:
        row=next(csv.reader([first],delimiter=delim))
    except Exception:
        row=[first]
    return {"type":"text_table","header":[x.strip() for x in row]}

def xlsx_schema(data):
    import openpyxl
    out={"type":"xlsx","sheets":[]}
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tf:
        tf.write(data); tf.flush()
        wb=openpyxl.load_workbook(tf.name,read_only=True,data_only=False)
        for ws in wb.worksheets:
            header=[]
            rows=ws.iter_rows(min_row=1,max_row=1,values_only=True)
            first=next(rows,())
            for v in first:
                header.append(None if v is None else str(v))
            out["sheets"].append({
                "name":ws.title,
                "max_row":ws.max_row,
                "max_column":ws.max_column,
                "header":header,
            })
    return out

def mat_schema(data):
    import scipy.io
    out={"type":"mat","variables":[]}
    with tempfile.NamedTemporaryFile(suffix=".mat") as tf:
        tf.write(data); tf.flush()
        for var,shape,klass in scipy.io.whosmat(tf.name):
            out["variables"].append({"name":var,"shape":list(shape),"class":klass})
    return out

def code_schema(data,ext):
    text=data.decode("utf-8","replace")
    # Strip comments approximately; collect identifiers/quoted strings only.
    if ext==".m":
        text_clean=re.sub(r"%[^\n]*","",text)
    elif ext in {".r",".R"}:
        text_clean=re.sub(r"#[^\n]*","",text)
    else:
        text_clean=text
    ids=sorted(set(re.findall(r"\b[A-Za-z][A-Za-z0-9_.]*\b",text_clean)))
    rel_ids=[x for x in ids if RELEVANT.search(x)][:1000]
    strings=sorted(set(re.findall(r'["\']([^"\'\n]{1,160})["\']',text_clean)))
    rel_strings=[x for x in strings if RELEVANT.search(x)][:1000]
    return {"type":"code","relevant_identifiers":rel_ids,"relevant_strings":rel_strings}

def inspect_member(z,info):
    name=info.filename
    ext=Path(name).suffix.lower()
    row={"name":name,"size":info.file_size,"extension":ext,"schema":None}
    if info.is_dir():
        row["schema"]={"type":"directory"}
        return row
    if info.file_size>100_000_000:
        row["schema"]={"type":"skipped_large"}
        return row
    if ext not in {".csv",".tsv",".txt",".xlsx",".xlsm",".mat",".m",".r",".py"}:
        row["schema"]={"type":"uninspected"}
        return row
    data=z.read(info)
    try:
        if ext in {".csv",".tsv",".txt"}:
            row["schema"]=csv_header(data,name)
        elif ext in {".xlsx",".xlsm"}:
            row["schema"]=xlsx_schema(data)
        elif ext==".mat":
            row["schema"]=mat_schema(data)
        elif ext in {".m",".r",".py"}:
            row["schema"]=code_schema(data,ext)
    except Exception as e:
        row["schema"]={"type":"inspection_error","error":repr(e)}
    return row

def main():
    blob=fetch_bytes(ZIP_URL)
    if not zipfile.is_zipfile(io.BytesIO(blob)):
        raise SystemExit("STOP: verified ZIP endpoint did not return ZIP")

    rows=[]
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        for info in z.infolist():
            rows.append(inspect_member(z,info))

    result={
        "source":ZIP_URL,
        "zip_bytes":len(blob),
        "member_count":len(rows),
        "members":rows,
    }
    OUT.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")

    lines=[
        "# Early-experience public schema audit v1","",
        "**SCHEMA ONLY — NO RESEARCH DATA ROWS / NUMERIC OUTCOMES REPORTED.**","",
        f"- ZIP bytes: {len(blob)}",
        f"- archive members: {len(rows)}","",
    ]
    for r in rows:
        lines += [f"## {r['name']}","",f"- bytes: {r['size']}"]
        s=r["schema"] or {}
        typ=s.get("type")
        lines.append(f"- schema type: {typ}")
        if typ=="text_table":
            lines.append(f"- header: {s.get('header')}")
        elif typ=="xlsx":
            for sh in s.get("sheets",[]):
                lines.append(
                    f"- sheet `{sh['name']}`: rows={sh['max_row']} "
                    f"cols={sh['max_column']} header={sh['header']}"
                )
        elif typ=="mat":
            for v in s.get("variables",[]):
                lines.append(f"- MAT `{v['name']}`: shape={v['shape']} class={v['class']}")
        elif typ=="code":
            lines.append(f"- relevant identifiers: {s.get('relevant_identifiers')}")
            lines.append(f"- relevant strings: {s.get('relevant_strings')}")
        elif typ=="inspection_error":
            lines.append(f"- error: {s.get('error')}")
        lines.append("")

    OUTMD.write_text("\n".join(lines)+"\n")
    print(OUTMD.read_text())

if __name__=="__main__":
    main()
