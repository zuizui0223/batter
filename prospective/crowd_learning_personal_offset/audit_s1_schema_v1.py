#!/usr/bin/env python3
from __future__ import annotations
from io import BytesIO
from pathlib import Path
import json
import urllib.request
import openpyxl

HERE=Path(__file__).resolve().parent
OUT=HERE/"S1_DATA_SCHEMA_V1.json"
OUTMD=HERE/"S1_DATA_SCHEMA_V1.md"

URLS=[
    "https://journals.plos.org/plosbiology/article/file?type=supplementary&id=info:doi/10.1371/journal.pbio.2002556.s011",
    "https://doi.org/10.1371/journal.pbio.2002556.s011",
]
HEAD={"User-Agent":"Mozilla/5.0 batter-crowd-learning-schema/1.0"}

def fetch():
    errors=[]
    for url in URLS:
        try:
            req=urllib.request.Request(url,headers=HEAD)
            with urllib.request.urlopen(req,timeout=180) as r:
                data=r.read()
            if data[:2]==b"PK":
                return data,url
            errors.append(f"{url}: not xlsx zip magic, bytes={len(data)}")
        except Exception as e:
            errors.append(f"{url}: {e!r}")
    raise RuntimeError(" ; ".join(errors))

def main():
    data,used=fetch()
    wb=openpyxl.load_workbook(BytesIO(data),read_only=False,data_only=False)
    sheets=[]
    for ws in wb.worksheets:
        strings=[]
        numeric_counts_by_row=[]
        max_scan=min(ws.max_row,25)
        for r in range(1,max_scan+1):
            nnum=0
            for c in range(1,ws.max_column+1):
                v=ws.cell(r,c).value
                if isinstance(v,(int,float)) and not isinstance(v,bool):
                    nnum+=1
                elif isinstance(v,str) and v.strip():
                    strings.append({"row":r,"col":c,"text":v.strip()[:300]})
            numeric_counts_by_row.append({"row":r,"n_numeric":nnum})
        col_numeric=[]
        for c in range(1,ws.max_column+1):
            n=0
            for r in range(1,ws.max_row+1):
                v=ws.cell(r,c).value
                if isinstance(v,(int,float)) and not isinstance(v,bool):
                    n+=1
            col_numeric.append({"col":c,"n_numeric":n})
        sheets.append({
            "name":ws.title,
            "max_row":ws.max_row,
            "max_column":ws.max_column,
            "merged_ranges":[str(x) for x in ws.merged_cells.ranges],
            "strings_first25":strings,
            "numeric_counts_first25_rows":numeric_counts_by_row,
            "numeric_counts_by_column":col_numeric,
        })
    wb.close()
    result={"source":used,"bytes":len(data),"sheets":sheets}
    OUT.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")
    lines=[
        "# Prat 2017 S1 Data structural audit v1","",
        "**STRUCTURE ONLY - NO NUMERIC RESEARCH VALUES REPORTED.**","",
        f"- source: {used}",
        f"- workbook bytes: {len(data)}",
        f"- sheets: {len(sheets)}",""
    ]
    for s in sheets:
        lines += [
            f"## {s['name']}","",
            f"- rows: {s['max_row']}",
            f"- columns: {s['max_column']}",
            f"- merged ranges: {s['merged_ranges']}",
            "- string labels in first 25 rows:"
        ]
        for q in s["strings_first25"]:
            lines.append(f"  - R{q['row']}C{q['col']}: {q['text']}")
        lines += ["- numeric cell counts by first 25 rows:"]
        for q in s["numeric_counts_first25_rows"]:
            lines.append(f"  - row {q['row']}: {q['n_numeric']}")
        lines.append("")
    OUTMD.write_text("\n".join(lines)+"\n")
    print(OUTMD.read_text())

if __name__=="__main__":
    main()
