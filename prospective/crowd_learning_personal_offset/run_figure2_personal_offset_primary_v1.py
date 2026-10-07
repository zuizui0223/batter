#!/usr/bin/env python3
"""Frozen Prat et al. 2017 Figure-2 within-group personal-offset primary."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import itertools
import json
import math
import re
import urllib.request

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
GATE = HERE / "FIGURE2_SUPPORT_RESULT_V1.json"
OUT = HERE / "FIGURE2_PERSONAL_OFFSET_RESULT_V1.json"
OUTMD = HERE / "FIGURE2_PERSONAL_OFFSET_RESULT_V1.md"

URL = "https://journals.plos.org/plosbiology/article/file?type=supplementary&id=info:doi/10.1371/journal.pbio.2002556.s011"
HEAD = {"User-Agent": "Mozilla/5.0 batter-crowd-learning-primary/1.0"}

GROUPS = ("High-F0", "Low-F0", "Control")
GROUP_BATS = {
    "High-F0": (1, 2, 3, 4),
    "Low-F0": (1, 2, 3, 4, 5),
    "Control": (1, 2, 3, 4, 5),
}
SESSIONS = (1, 2, 3, 4)
MIN_PAIRED = 5

PATTERN = re.compile(
    r"^session\s+(?P<session>[1-4]),"
    r"(?P<group>High-F0|Low-F0|Control),"
    r"bat\s+(?P<bat>[1-5])\s+\(LD(?P<axis>[12])\)$"
)


def fetch() -> bytes:
    req = urllib.request.Request(URL, headers=HEAD)
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = resp.read()
    if data[:2] != b"PK":
        raise RuntimeError("supplement is not an XLSX archive")
    return data


def finite_number(x) -> bool:
    if isinstance(x, bool):
        return False
    if not isinstance(x, (int, float)):
        return False
    return math.isfinite(float(x))


def load_states():
    wb = openpyxl.load_workbook(BytesIO(fetch()), read_only=True, data_only=True)
    ws = wb["Figure 2"]

    columns = {}
    for col in range(1, ws.max_column + 1):
        label = ws.cell(2, col).value
        if not isinstance(label, str):
            continue
        m = PATTERN.match(label.strip())
        if not m:
            continue
        key = (m.group("group"), int(m.group("bat")), int(m.group("session")))
        axis = int(m.group("axis"))
        columns.setdefault(key, {})[axis] = col

    states = {}
    support = {}
    for group in GROUPS:
        for bat in GROUP_BATS[group]:
            for session in SESSIONS:
                key = (group, bat, session)
                if key not in columns or set(columns[key]) != {1, 2}:
                    raise RuntimeError(f"missing LD pair: {key}")
                c1, c2 = columns[key][1], columns[key][2]
                pairs = []
                for row in range(3, ws.max_row + 1):
                    a = ws.cell(row, c1).value
                    b = ws.cell(row, c2).value
                    if finite_number(a) and finite_number(b):
                        pairs.append((float(a), float(b)))
                if len(pairs) < MIN_PAIRED:
                    raise RuntimeError(f"support drift {key}: {len(pairs)}")
                arr = np.asarray(pairs, dtype=float)
                states[key] = arr.mean(axis=0)
                support[f"{group}|bat{bat}|session{session}"] = int(len(arr))

    wb.close()
    if len(states) != 56:
        raise RuntimeError(f"expected 56 pup-session states, got {len(states)}")
    return states, support


def residualize_group_session(states):
    residual = {}
    for group in GROUPS:
        bats = GROUP_BATS[group]
        for session in SESSIONS:
            mat = np.vstack([states[(group, bat, session)] for bat in bats])
            mu = mat.mean(axis=0)
            for bat in bats:
                residual[(group, bat, session)] = states[(group, bat, session)] - mu
    return residual


def histories_and_targets(residual):
    histories = {}
    targets = {}
    for group in GROUPS:
        for bat in GROUP_BATS[group]:
            histories[(group, bat)] = np.mean(
                np.vstack([residual[(group, bat, s)] for s in (1, 2, 3)]),
                axis=0,
            )
            targets[(group, bat)] = residual[(group, bat, 4)]
    return histories, targets


def identity_advantages(histories, targets):
    values = {}
    for group in GROUPS:
        bats = GROUP_BATS[group]
        for bat in bats:
            target = targets[(group, bat)]
            self_d = float(np.linalg.norm(target - histories[(group, bat)]))
            donor_d = [
                float(np.linalg.norm(target - histories[(group, other)]))
                for other in bats
                if other != bat
            ]
            values[(group, bat)] = float(np.mean(donor_d) - self_d)
    return values


def group_perm_sums(group, histories, targets):
    """Return exact sum of K_i within a group for every target-label permutation."""
    bats = GROUP_BATS[group]
    out = []
    for perm in itertools.permutations(bats):
        # perm is target-source bat aligned to the fixed history identity order.
        total = 0.0
        for identity, target_source in zip(bats, perm):
            target = targets[(group, target_source)]
            self_d = float(np.linalg.norm(target - histories[(group, identity)]))
            donor_d = [
                float(np.linalg.norm(target - histories[(group, other)]))
                for other in bats
                if other != identity
            ]
            total += float(np.mean(donor_d) - self_d)
        out.append(total)
    return np.asarray(out, dtype=float)


def main():
    gate = json.loads(GATE.read_text())
    if gate.get("gate") != "PASS_FIGURE2_REPEATED_SUPPORT":
        raise SystemExit("STOP: repeated-support gate did not authorize numeric opening")
    if gate.get("reconstructed_units") != 56:
        raise SystemExit("STOP: gate unit count drift")
    if gate.get("minimum_paired_support", 0) < MIN_PAIRED:
        raise SystemExit("STOP: gate support drift")

    states, support = load_states()
    residual = residualize_group_session(states)
    histories, targets = histories_and_targets(residual)

    observed = identity_advantages(histories, targets)
    obs_sum = float(sum(observed.values()))
    obs_k = obs_sum / 14.0

    bat_rows = []
    group_mean = {}
    for group in GROUPS:
        vals = []
        for bat in GROUP_BATS[group]:
            v = observed[(group, bat)]
            vals.append(v)
            bat_rows.append({"group": group, "bat": bat, "K_i": v})
        group_mean[group] = float(np.mean(vals))

    # Exact null factorizes by group. Combine the exact group permutation sums.
    high = group_perm_sums("High-F0", histories, targets)  # 4! = 24
    low = group_perm_sums("Low-F0", histories, targets)    # 5! = 120
    ctrl = group_perm_sums("Control", histories, targets)  # 5! = 120

    if (len(high), len(low), len(ctrl)) != (24, 120, 120):
        raise RuntimeError("unexpected permutation counts")

    extreme = 0
    greater = 0
    n_null = 0
    tol = 1e-15
    null_sum = 0.0
    null_sq = 0.0

    for h in high:
        for l in low:
            vals = (h + l + ctrl) / 14.0
            n_null += len(vals)
            extreme += int(np.sum(vals >= obs_k - tol))
            greater += int(np.sum(vals > obs_k + tol))
            null_sum += float(np.sum(vals))
            null_sq += float(np.sum(vals * vals))

    if n_null != 345600:
        raise RuntimeError(f"expected 345600 null assignments, got {n_null}")

    p_exact = extreme / n_null
    rank_desc = 1 + greater
    null_mean = null_sum / n_null
    null_sd = max(0.0, null_sq / n_null - null_mean * null_mean) ** 0.5

    loo = {}
    for row in bat_rows:
        key = f"{row['group']}|bat{row['bat']}"
        remaining = [
            x["K_i"]
            for x in bat_rows
            if not (x["group"] == row["group"] and x["bat"] == row["bat"])
        ]
        loo[key] = float(np.mean(remaining))

    positive = sum(row["K_i"] > 0 for row in bat_rows)
    verdict = (
        "SUPPORTED"
        if obs_k > 0 and p_exact <= 0.05
        else ("POSITIVE_BUT_UNRESOLVED" if obs_k > 0 else "NO_POSITIVE_IDENTITY")
    )

    result = {
        "version": 1,
        "source": "10.1371/journal.pbio.2002556.s011",
        "figure": 2,
        "groups": list(GROUPS),
        "n_pups": 14,
        "sessions": list(SESSIONS),
        "minimum_paired_support": min(support.values()),
        "support": support,
        "K": obs_k,
        "positive_pups": positive,
        "p_exact": p_exact,
        "rank_descending": rank_desc,
        "n_exact_assignments": n_null,
        "null_mean": null_mean,
        "null_sd": null_sd,
        "bat_advantages": bat_rows,
        "group_mean_advantage": group_mean,
        "leave_one_pup_out_K_descriptive": loo,
        "verdict": verdict,
    }
    OUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Prat 2017 Figure-2 personal-offset result v1",
        "",
        f"- pups = **14**",
        f"- minimum paired calls per pup × session = **{result['minimum_paired_support']}**",
        f"- K = **{obs_k:+.6f}**",
        f"- positive pups = **{positive}/14**",
        f"- exact assignments = **{n_null:,}**",
        f"- exact one-sided p = **{p_exact:.8f}**",
        f"- observed rank = **{rank_desc} / {n_null:,}**",
        f"- verdict = **{verdict}**",
        "",
        "## Group mean advantages — descriptive only",
        "",
    ]
    for group in GROUPS:
        lines.append(f"- {group}: {group_mean[group]:+.6f}")

    lines += ["", "## Pup-level advantages", ""]
    for row in bat_rows:
        lines.append(f"- {row['group']} bat {row['bat']}: {row['K_i']:+.6f}")

    lines += [
        "",
        "The null independently permutes session-4 identities within each source group.",
        "Calls are measurement replicates; all 14 pups receive equal programme weight.",
        "No playback-treatment causal claim is made from this identity test.",
        "",
    ]
    OUTMD.write_text("\n".join(lines) + "\n")
    print(OUTMD.read_text())


if __name__ == "__main__":
    main()
