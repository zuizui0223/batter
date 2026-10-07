#!/usr/bin/env python3
"""Structure-only paired-call support audit for Prat et al. 2017 Figure 2.

No LD1/LD2 coordinate magnitudes are printed or summarized.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import json
import math
import re
import urllib.request

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE / "FIGURE2_SUPPORT_RESULT_V1.json"
OUTMD = HERE / "FIGURE2_SUPPORT_RESULT_V1.md"

URL = "https://journals.plos.org/plosbiology/article/file?type=supplementary&id=info:doi/10.1371/journal.pbio.2002556.s011"
HEAD = {"User-Agent": "Mozilla/5.0 batter-crowd-learning-support/1.0"}

GROUP_ORDER = ("High-F0", "Low-F0", "Control")
EXPECTED_BATS = {"High-F0": 4, "Low-F0": 5, "Control": 5}
EXPECTED_SESSIONS = (1, 2, 3, 4)
MIN_PAIRED = 5

PATTERN = re.compile(
    r"^session\s+(?P<session>[1-4]),"
    r"(?P<group>High-F0|Low-F0|Control),"
    r"bat\s+(?P<bat>[1-5])\s+\(LD(?P<axis>[12])\)$"
)


def fetch() -> bytes:
    req = urllib.request.Request(URL, headers=HEAD)
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = resp.read()
    if data[:2] != b"PK":
        raise RuntimeError("supplement is not an XLSX zip archive")
    return data


def finite_number(x) -> bool:
    if isinstance(x, bool):
        return False
    if not isinstance(x, (int, float)):
        return False
    return math.isfinite(float(x))


def key_text(session: int, group: str, bat: int) -> str:
    return f"session {session}|{group}|bat {bat}"


def main() -> None:
    wb = openpyxl.load_workbook(BytesIO(fetch()), read_only=True, data_only=True)
    ws = wb["Figure 2"]

    columns = {}
    for col in range(1, ws.max_column + 1):
        label = ws.cell(2, col).value
        if not isinstance(label, str):
            continue
        m = PATTERN.match(label.strip())
        if not m:
            continue
        session = int(m.group("session"))
        group = m.group("group")
        bat = int(m.group("bat"))
        axis = int(m.group("axis"))
        key = (session, group, bat)
        columns.setdefault(key, {})
        if axis in columns[key]:
            raise RuntimeError(f"duplicate LD{axis} column for {key}")
        columns[key][axis] = col

    expected = []
    for session in EXPECTED_SESSIONS:
        for group in GROUP_ORDER:
            for bat in range(1, EXPECTED_BATS[group] + 1):
                expected.append((session, group, bat))

    missing_units = [key_text(*k) for k in expected if k not in columns]
    unexpected_units = [key_text(*k) for k in columns if k not in set(expected)]
    incomplete_pairs = [
        key_text(*k)
        for k in expected
        if k in columns and set(columns[k]) != {1, 2}
    ]

    records = []
    for session, group, bat in expected:
        key = (session, group, bat)
        if key not in columns or set(columns[key]) != {1, 2}:
            continue

        c1 = columns[key][1]
        c2 = columns[key][2]
        n1 = 0
        n2 = 0
        npaired = 0
        for row in range(3, ws.max_row + 1):
            v1 = ws.cell(row, c1).value
            v2 = ws.cell(row, c2).value
            f1 = finite_number(v1)
            f2 = finite_number(v2)
            n1 += int(f1)
            n2 += int(f2)
            npaired += int(f1 and f2)

        records.append(
            {
                "session": session,
                "group": group,
                "bat": bat,
                "ld1_numeric": n1,
                "ld2_numeric": n2,
                "paired_finite": npaired,
                "pass_min5": npaired >= MIN_PAIRED,
            }
        )

    wb.close()

    all_expected = (
        len(columns) == 56
        and not missing_units
        and not unexpected_units
        and not incomplete_pairs
        and len(records) == 56
    )
    support_pass = all(r["pass_min5"] for r in records)
    gate = (
        "PASS_FIGURE2_REPEATED_SUPPORT"
        if all_expected and support_pass
        else "STOP_FIGURE2_REPEATED_SUPPORT"
    )

    result = {
        "version": 1,
        "source": "10.1371/journal.pbio.2002556.s011",
        "sheet": "Figure 2",
        "min_paired_calls": MIN_PAIRED,
        "expected_units": 56,
        "reconstructed_units": len(records),
        "missing_units": missing_units,
        "unexpected_units": unexpected_units,
        "incomplete_pairs": incomplete_pairs,
        "minimum_paired_support": min(
            (r["paired_finite"] for r in records), default=0
        ),
        "records": records,
        "gate": gate,
    }
    OUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Prat 2017 Figure-2 support result v1",
        "",
        "**STRUCTURAL COUNTS ONLY — NO LD COORDINATE MAGNITUDES REPORTED.**",
        "",
        f"- reconstructed pup × session units: **{len(records)} / 56**",
        f"- minimum paired LD1/LD2 call support: **{result['minimum_paired_support']}**",
        f"- frozen support floor: **{MIN_PAIRED}**",
        f"- gate: **{gate}**",
        "",
        "| session | group | bat | LD1 numeric | LD2 numeric | paired finite | pass |",
        "|---:|---|---:|---:|---:|---:|---|",
    ]
    for r in records:
        lines.append(
            f"| {r['session']} | {r['group']} | {r['bat']} | "
            f"{r['ld1_numeric']} | {r['ld2_numeric']} | "
            f"{r['paired_finite']} | {'PASS' if r['pass_min5'] else 'FAIL'} |"
        )

    if missing_units:
        lines += ["", "## Missing units", ""]
        lines += [f"- {x}" for x in missing_units]
    if incomplete_pairs:
        lines += ["", "## Incomplete LD pairs", ""]
        lines += [f"- {x}" for x in incomplete_pairs]

    lines += [
        "",
        "Zero-valued LDA coordinates were treated as valid finite coordinates.",
        "No LD magnitude, centroid, distance or identity statistic was calculated.",
        "",
    ]
    OUTMD.write_text("\n".join(lines) + "\n")
    print(OUTMD.read_text())


if __name__ == "__main__":
    main()
