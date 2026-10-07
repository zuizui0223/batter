#!/usr/bin/env python3
"""Schema-only inspection of public perturbation datasets.

Downloads public research files but extracts structural metadata only.
Never prints or computes numeric research outcomes.
"""

from __future__ import annotations

import io
import json
import pathlib
import re
import tempfile
import urllib.parse
import urllib.request
import zipfile

HERE = pathlib.Path(__file__).resolve().parent
OUT_JSON = HERE / "PUBLIC_SCHEMA_AUDIT_RESULT_V1.json"
OUT_MD = HERE / "PUBLIC_SCHEMA_AUDIT_RESULT_V1.md"

SOURCES = [
    {"key":"aharon2017","dataset_id":"f6mvhj5gj9","version":3},
    {"key":"ma2025","dataset_id":"964fv73w94","version":1},
]

HEADERS = {
    "User-Agent":"batter-schema-only-audit/1.0",
    "Accept":"application/json, application/octet-stream, */*",
}


def req(url):
    return urllib.request.Request(url, headers=HEADERS)


def get_json(url):
    with urllib.request.urlopen(req(url), timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))


def list_files(dataset_id, version):
    url = (
        f"https://data.mendeley.com/api/datasets/{dataset_id}/files?"
        + urllib.parse.urlencode({"version":version,"$start":0,"$limit":1000})
    )
    data = get_json(url)
    if isinstance(data, dict):
        return data.get("items") or data.get("files") or data.get("results") or data.get("data") or []
    return data or []


def download_bytes(url):
    with urllib.request.urlopen(req(url), timeout=120) as r:
        return r.read()


def get_download_url(dataset_id, version, row):
    details = row.get("content_details") or {}
    return details.get("download_url") or row.get("download_url")


def safe_name(name):
    return pathlib.Path(name or "unnamed").name


def inspect_mat_bytes(data, name):
    out={"type":"mat","variables":[],"notes":[]}
    # v7.3 MAT files are HDF5.
    if data[:8] == b"\x89HDF\r\n\x1a\n":
        import h5py
        with h5py.File(io.BytesIO(data),"r") as h:
            def visit(path,obj):
                if isinstance(obj,h5py.Dataset):
                    out["variables"].append({
                        "name":path,
                        "shape":list(obj.shape),
                        "dtype":str(obj.dtype),
                    })
            h.visititems(visit)
        return out
    try:
        import scipy.io
        with tempfile.NamedTemporaryFile(suffix=".mat") as tf:
            tf.write(data); tf.flush()
            for var,shape,klass in scipy.io.whosmat(tf.name):
                out["variables"].append({
                    "name":var,
                    "shape":list(shape),
                    "class":klass,
                })
    except Exception as exc:
        out["notes"].append("whosmat_failed:"+repr(exc))
    return out


def inspect_xlsx_bytes(data, name):
    import openpyxl
    out={"type":"xlsx","sheets":[]}
    wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
    for ws in wb.worksheets:
        first=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
        out["sheets"].append({
            "name":ws.title,
            "max_row":ws.max_row,
            "max_column":ws.max_column,
            "header":[None if v is None else str(v) for v in first],
        })
    wb.close()
    return out


def inspect_matlab_script(data, name):
    text=data.decode("utf-8","replace")
    # Strip comments before token extraction.
    text_no_comments=re.sub(r"%[^\n]*","",text)
    identifiers=sorted(set(re.findall(r"\b[A-Za-z]\w*\b",text_no_comments)))
    relevant=[
        x for x in identifiers
        if re.search(r"(bat|trial|fly|flight|noise|prey|land|speed|path|turn|condition|xyz|coord)",x,re.I)
    ][:500]
    calls=[]
    pat=r"\b(load|readtable|readmatrix|readcell|xlsread)\s*\(\s*['\"]([^'\"]+)['\"]"
    for func,filename in re.findall(pat,text_no_comments,re.I):
        calls.append({"function":func,"filename":filename})
    strings=sorted(set(
        s for s in re.findall(r"['\"]([^'\"\n]{1,160})['\"]",text_no_comments)
        if re.search(r"(bat|trial|noise|prey|land|fly|\.mat|\.csv|\.xlsx|data)",s,re.I)
    ))[:300]
    return {
        "type":"matlab_script",
        "relevant_identifiers":relevant,
        "data_file_references":calls,
        "relevant_strings":strings,
    }


def inspect_zip_bytes(data,name):
    out={"type":"zip","members":[]}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist()[:5000]:
            info=z.getinfo(n)
            out["members"].append({"name":n,"size":info.file_size})
    return out


def inspect_one(dataset_id,version,row):
    name=row.get("filename") or row.get("name") or ""
    ext=pathlib.Path(name).suffix.lower()
    allowed={".mat",".m",".xlsx",".xlsm",".zip",".csv",".txt"}
    result={
        "filename":name,
        "id":row.get("id"),
        "extension":ext,
        "inspection":None,
        "status":None,
    }
    if ext not in allowed:
        result["status"]="SKIP_UNSUPPORTED_STRUCTURE_ONLY_TYPE"
        return result
    url=get_download_url(dataset_id,version,row)
    if not url:
        result["status"]="STOP_NO_DOWNLOAD_URL"
        return result
    data=download_bytes(url)
    if ext==".mat":
        result["inspection"]=inspect_mat_bytes(data,name)
    elif ext==".m":
        result["inspection"]=inspect_matlab_script(data,name)
    elif ext in {".xlsx",".xlsm"}:
        result["inspection"]=inspect_xlsx_bytes(data,name)
    elif ext==".zip":
        result["inspection"]=inspect_zip_bytes(data,name)
    elif ext in {".csv",".txt"}:
        # Header only.
        text=data.decode("utf-8","replace")
        first=text.splitlines()[0] if text.splitlines() else ""
        result["inspection"]={"type":"text_header_only","header":first[:2000]}
    result["status"]="STRUCTURE_INSPECTED"
    return result


def derive_tokens(source):
    names=[]
    for f in source["files"]:
        ins=f.get("inspection") or {}
        for v in ins.get("variables",[]):
            names.append(v.get("name",""))
        names.extend(ins.get("relevant_identifiers",[]))
        names.extend(x.get("filename","") for x in ins.get("data_file_references",[]))
        names.extend(ins.get("relevant_strings",[]))
        names.extend(x.get("name","") for x in ins.get("members",[]))
    bat_tokens=sorted(set(re.findall(r"(?i)\b(?:bat)?\d{2,4}\b"," ".join(names))))
    condition_like=sorted(set(x for x in names if re.search(r"(con|cond|wind|slow|turn|noise|prey|land)",x,re.I)))[:300]
    return {"bat_tokens":bat_tokens[:200],"condition_like_tokens":condition_like}


def render_md(result):
    lines=["# Public perturbation schema audit result v1","","## Status","",
           "**SCHEMA-ONLY OPENING; NO NUMERIC OUTCOME VALUES REPORTED.**",""]
    for s in result["sources"]:
        lines += [f"## {s['key']}","",f"- file count: {len(s['files'])}"]
        toks=s.get("tokens") or {}
        lines += [
            f"- bat-like tokens: {', '.join(toks.get('bat_tokens') or []) or 'none'}",
            f"- condition-like identifiers: {len(toks.get('condition_like_tokens') or [])}",
            "",
        ]
        for f in s["files"]:
            lines.append(f"### {f['filename']}")
            lines.append(f"- status: {f['status']}")
            ins=f.get("inspection") or {}
            if ins.get("type")=="mat":
                lines.append(f"- MAT variables: {len(ins.get('variables') or [])}")
                for v in (ins.get("variables") or [])[:120]:
                    lines.append(f"  - `{v.get('name')}` shape={v.get('shape')} class={v.get('class') or v.get('dtype')}")
            elif ins.get("type")=="matlab_script":
                lines.append(f"- relevant identifiers: {', '.join(ins.get('relevant_identifiers') or [])}")
                for x in ins.get("data_file_references") or []:
                    lines.append(f"  - data reference: {x}")
            elif ins.get("type")=="xlsx":
                for sh in ins.get("sheets") or []:
                    lines.append(f"  - sheet `{sh['name']}`: rows={sh['max_row']} cols={sh['max_column']} header={sh['header']}")
            elif ins.get("type")=="zip":
                lines.append(f"- zip members: {len(ins.get('members') or [])}")
                for x in (ins.get("members") or [])[:120]:
                    lines.append(f"  - {x['name']} ({x['size']} bytes)")
            lines.append("")
    lines += ["## Boundary","","Shapes/counts/tokens are structural support only; no biological outcome was calculated.",""]
    return "\n".join(lines)


def main():
    out={"audit_version":1,"sources":[]}
    for src in SOURCES:
        row={"key":src["key"],"dataset_id":src["dataset_id"],"version":src["version"],"files":[]}
        try:
            files=list_files(src["dataset_id"],src["version"])
            for f in files:
                if isinstance(f,dict):
                    try:
                        row["files"].append(inspect_one(src["dataset_id"],src["version"],f))
                    except Exception as exc:
                        row["files"].append({
                            "filename":f.get("filename") or f.get("name"),
                            "id":f.get("id"),
                            "status":"INSPECTION_ERROR",
                            "error":repr(exc),
                        })
            row["tokens"]=derive_tokens(row)
        except Exception as exc:
            row["error"]=repr(exc)
        out["sources"].append(row)
    OUT_JSON.write_text(json.dumps(out,indent=2,ensure_ascii=False)+"\n")
    OUT_MD.write_text(render_md(out)+"\n")
    print(OUT_MD.read_text())


if __name__=="__main__":
    main()
